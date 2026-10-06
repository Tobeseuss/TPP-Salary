# -*- coding: utf-8 -*-
"""تست‌های نسخه 1.8.0 — درخواست کاربر:

«خروجی pdf در نسخه پایتون بسیار افتضاح هست و به بازنویسی کامل دارد …
قسمت دریافت فیش های حقوقی و فیش بانک در نسخه پایتون وجود ندارد …
قابلیت های پشتیبان گیری و بازگردانی پشتیبان در نسخه پایتون تعریف نشده …
در هر دو نسخه پلاگین و پایتون میخواهم بخشی به عنوان گزارش سالانه مراکز
داشته باشم که گزارش لیست حقوق یک مرکز را در ماه های مختلف یک سال در قالب
چندین شیت درون یک فایل اکسل خروجی می دهد»

پوشش تست (برنامه پایتون):
  ۱) موتور PDF جدید — سه قالب (گزارش/فیش بانکی/فیش A5) + چندصفحه‌ای
  ۲) صفحه فیش‌های حقوقی — فهرست + ساخت PDF فیش + فیلترهای چاپی
  ۳) صفحه فیش بانکی — سطرها از پروفایل + اکسل + PDF
  ۴) گزارش سالانه مراکز — آمار ماهانه + اکسل چندشیتی (جمع سال + شیت ماه‌ها)
  ۵) پشتیبان JSON/ZIP + بازگردانی کامل + بازگردانی بکاپ افزونه (سازگاری)
Run: QT_QPA_PLATFORM=offscreen python3 scripts/test_pyapp_180.py
"""
import os
import sys
import tempfile
import traceback

os.environ["QT_QPA_PLATFORM"] = "offscreen"
APP = os.environ.get("TPP_SMOKE_APP", "/home/z/my-project/build/tpp_salary/python-app")
sys.path.insert(0, APP)
sys.path.insert(0, "/home/z/my-project/scripts")

fail = 0
def check(cond, msg):
    global fail
    print(("  PASS: " if cond else "  FAIL: ") + msg)
    if not cond:
        fail = 1

from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

# stub modals
QMessageBox.information = staticmethod(lambda *a, **k: QMessageBox.StandardButton.Ok)
QMessageBox.warning = staticmethod(lambda *a, **k: QMessageBox.StandardButton.Ok)
QMessageBox.question = staticmethod(lambda *a, **k: QMessageBox.StandardButton.Yes)
QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: ("", ""))
QFileDialog.getOpenFileName = staticmethod(lambda *a, **k: ("", ""))

app = QApplication.instance() or QApplication([])
tmp = tempfile.mkdtemp(prefix="tpp180_")

from app.ui.app_window import MainWindow  # noqa: E402
from app import pdf_engine as PE          # noqa: E402
from app.api_client import load_json      # noqa: E402

win = MainWindow(tmp)

print("== 1) موتور PDF جدید ==")
out1 = os.path.join(tmp, "rep.pdf")
emps = [(1, "علی رضایی"), (2, "مریم احمدی"), (3, "حسین کریمی"), (4, "زهرا موسوی"), (5, "رضا محمدی")]
rows = [
    ("حقوق پایه", {1: 250000000, 2: 180000000, 3: 210000000, 4: 195000000, 5: 230000000}, False),
    ("جمع ناخالص", {1: 262000000, 2: 192000000, 3: 222000000, 4: 207000000, 5: 242000000}, True),
]
PE.write_report_pdf(out1, emps, rows, "شرکت آزمون", ["دوره: مرداد 1404 — مرکز: مرکز اصلی"])
check(os.path.getsize(out1) > 3000, "PDF گزارش ساخته شد (%d بایت)" % os.path.getsize(out1))
many = [("سطر %d" % i, {e[0]: 1000000 * i for e in emps}, i % 5 == 0) for i in range(1, 61)]
out1b = os.path.join(tmp, "rep_multi.pdf")
pages = PE.write_report_pdf(out1b, emps, many, "شرکت آزمون", ["دوره: مرداد 1404"])
check(pages > 1 and os.path.getsize(out1b) > 10000, "PDF چندصفحه‌ای با تکرار سربرگ (%d صفحه)" % pages)
out1c = os.path.join(tmp, "bank.pdf")
PE.write_bank_pdf(out1c, [("علی رضایی", "1234567890", "IR820540102680020817909002", 210000000)],
                  "فیش بانکی بانک ملت — مرداد 1404", ["کارکنان دارای حساب: 1 نفر"])
check(os.path.getsize(out1c) > 3000, "PDF فیش بانکی ساخته شد")
out1d = os.path.join(tmp, "payslip.pdf")
PE.write_payslip_pdf(out1d, {"name": "علی رضایی", "national": "0012345678", "job_title": "کارشناس",
                             "personnel": 1, "period_label": "مرداد 1404", "center_name": "مرکز اصلی"},
                     [("حقوق پایه", "250,000,000", True, 250000000.0),
                      ("کسر بیمه", "-18,948,842", True, -18948842.0)],
                     "شرکت آزمون", "ریال", "1404/05/12")
check(os.path.getsize(out1d) > 3000, "PDF فیش حقوقی A5 ساخته شد")
check(PE.ensure_fonts() in ("Vazirmatn", "Tahoma", "Segoe UI", "DejaVu Sans"),
      "فونت برنامه بارگذاری شد: %s" % PE.ensure_fonts())

print("== 2) صفحه فیش‌های حقوقی ==")
# داده آزمون: مرکز + دو کارمند + رکورد
cen_id = win.store.save_center(0, "مرکز آزمون ۱۸۰")
e1 = win.store.save_employee({"id": 0, "name": "علی رضایی", "national": "1111111111", "centers": [cen_id]})
e2 = win.store.save_employee({"id": 0, "name": "مریم احمدی", "national": "2222222222", "centers": [cen_id]})
r1 = win.store.upsert_record(e1, cen_id, 1404, 5, {"base_salary": 250000000, "gross": 250000000,
                                                   "insurable": 250000000, "insurance_deduct": 18948842,
                                                   "other_deductions": 0, "net": 231051158})
r2 = win.store.upsert_record(e2, cen_id, 1404, 5, {"base_salary": 180000000, "gross": 180000000,
                                                   "insurable": 180000000, "insurance_deduct": 13662000,
                                                   "other_deductions": 0, "net": 166338000})
ps = win.pages["payslips"]
ps.cmb_year.setCurrentIndex(ps.cmb_year.findData(1404))
ps.cmb_month.setCurrentIndex(4)  # مرداد
idx = ps.cmb_center.findData(cen_id)
ps.cmb_center.setCurrentIndex(idx)
ps.load_rows()
check(ps.table.rowCount() == 2, "فهرست فیش‌ها = 2 (دو رکورد دوره)")
check(ps.table.cellWidget(0, 4) is not None, "دکمه مشاهده/چاپ در هر ردیف")

pdf_tmp = os.path.join(tmp, "one.pdf")
ok, err = ps._make_pdf(r1, pdf_tmp)
check(ok and os.path.exists(pdf_tmp) and os.path.getsize(pdf_tmp) > 3000,
      "ساخت PDF فیش تکی موفق (%s)" % (err or "%d بایت" % os.path.getsize(pdf_tmp)))
# فیلدهای چاپی — با ساختار فیلدهای واقعی (هم‌سان bundle سرور)
for fk, fl, ft in (("base_salary", "حقوق پایه", "number"),
                   ("overtime_hours", "ساعات اضافه‌کاری", "number"),
                   ("insurance_deduct", "بیمه سهم کارمند", "number"),
                   ("insurance_group", "گروه اصلی بیمه", "text")):
    win.db.ex("INSERT OR REPLACE INTO fields(field_key, label, type, default_value, formula, calculated, allow_manual, sort)"
              " VALUES(?,?,?,?,?,0,1,?)", (fk, fl, ft, "0", "", 1))
from app.ui.pages.payslips import payslip_fields
fields = payslip_fields(win.store, {"base_salary": 250000000, "overtime_hours": 40,
                                    "insurance_deduct": 0, "insurance_group": "گروه ۳"})
labels = [f[0] for f in fields]
check(not any("اضافه" in l for l in labels),
      "فیلدهای فقط‌محاسباتی (ساعات اضافه‌کاری) در فیش چاپ نمی‌شوند")
check(not any(f[0] == "بیمه سهم کارمند" for f in fields), "فیلد عددی صفر (بیمه) در فیش چاپ نمی‌شود")
check(any(f[0] == "گروه اصلی بیمه" for f in fields), "فیلد متنی غیرخالی چاپ می‌شود")
check(any(f[0] == "حقوق پایه" for f in fields), "فیلد عددی ناصفر چاپ می‌شود")

# ZIP عمده — با دیالوگ استاب‌شده به مسیر واقعی
zip_path = os.path.join(tmp, "payslips.zip")
QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (zip_path, ""))
ps.btn_zip.click()
import zipfile as _zip
check(os.path.exists(zip_path) and os.path.getsize(zip_path) > 3000, "ZIP عمده فیش‌ها ساخته شد")
with _zip.ZipFile(zip_path) as z:
    check(len(z.namelist()) == 2, "دو فیش PDF داخل ZIP: %s" % z.namelist())
QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: ("", ""))

print("== 3) صفحه فیش بانکی ==")
# بانک + حساب در پروفایل کارمند
b1 = win.store.save_bank(0, "بانک آزمون")
e1d = win.store.employee(e1)
prof = dict(e1d.get("profile") or {})
prof["bank_accounts"] = {str(b1): {"account": "1234567890", "sheba": "IR820540102680020817909002", "card": ""}}
win.store.save_employee({"id": e1, "name": "علی رضایی", "national": "1111111111",
                         "centers": [cen_id], "profile": prof, "terminated": False})
e2d = win.store.employee(e2)
prof2 = dict(e2d.get("profile") or {})
prof2["bank_accounts"] = {str(b1): {"account": "", "sheba": "", "card": ""}}  # بدون شماره حساب → حذف از فیش
win.store.save_employee({"id": e2, "name": "مریم احمدی", "national": "2222222222",
                         "centers": [cen_id], "profile": prof2, "terminated": False})
bf = win.pages["bankfiche"]
bf.refresh()
bf.cmb_year.setCurrentIndex(bf.cmb_year.findData(1404))
bf.cmb_month.setCurrentIndex(4)
bf.cmb_center.setCurrentIndex(bf.cmb_center.findData(cen_id))
bf.cmb_bank.setCurrentIndex(bf.cmb_bank.findData(b1))
bf.load_rows()
check(len(bf._rows) == 1, "فیش بانکی فقط کارمند دارای شماره حساب (1 سطر)")
check(bf._rows[0][0] == "علی رضایی" and bf._rows[0][1] == "1234567890", "نام و شماره حساب درست")
check(abs(bf._rows[0][3] - 231051158) < 1, "خالص پرداختی از رکورد دوره")
# اکسل فیش بانکی
xlsx1 = os.path.join(tmp, "bank.xlsx")
QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (xlsx1, ""))
bf.export_excel()
from openpyxl import load_workbook
wb1 = load_workbook(xlsx1)
ws1 = wb1.active
check(ws1.title == "فیش بانکی" and ws1.cell(row=3, column=1).value == "#", "اکسل فیش بانکی: سربرگ و ستون‌ها")
check(ws1.cell(row=4, column=2).value == "علی رضایی", "اکسل فیش بانکی: سطر کارمند")
pdf_bank = os.path.join(tmp, "bank_out.pdf")
QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (pdf_bank, ""))
bf.export_pdf()
check(os.path.exists(pdf_bank) and os.path.getsize(pdf_bank) > 3000, "خروجی PDF فیش بانکی از صفحه")
QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: ("", ""))

print("== 4) گزارش سالانه مراکز ==")
# ماه دوم (شهریور) هم اضافه کن
r3 = win.store.upsert_record(e1, cen_id, 1404, 6, {"base_salary": 260000000, "gross": 260000000,
                                                   "insurable": 260000000, "insurance_deduct": 19707000,
                                                   "other_deductions": 0, "net": 240293000})
YEARS_ = __import__("app.ui.pages.annual", fromlist=["YEARS"]).YEARS
an = win.pages["annual"]
an.refresh()
an.cmb_year.setCurrentIndex(YEARS_.index(1404))  # cmb_year بدون data role — ایندکسی
an.cmb_center.setCurrentIndex(an.cmb_center.findData(cen_id))
an.build()
check(len(an._stats) == 2, "دو ماه دارای رکورد (مرداد/شهریور)")
check(an._stats[0]["jmonth"] == 5 and an._stats[1]["jmonth"] == 6, "ماه‌ها مرتب")
check(abs(an._stats[0]["net"] - (231051158 + 166338000)) < 1, "جمع خالص مرداد = دو رکورد")
xlsx2 = os.path.join(tmp, "annual.xlsx")
QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (xlsx2, ""))
an.export_excel()
wb2 = load_workbook(xlsx2)
check("جمع سال" in wb2.sheetnames, "شیت «جمع سال» موجود")
check("مرداد" in wb2.sheetnames and "شهریور" in wb2.sheetnames, "شیت ماه‌ها موجود: %s" % wb2.sheetnames)
wsm = wb2["مرداد"]
check(wsm.cell(row=4, column=1).value == "عناوین", "قالب ستونی شیت ماه (سطر ۴ عناوین)")
hdrs = [wsm.cell(row=4, column=c).value for c in range(2, wsm.max_column + 1)]
check("علی رضایی" in hdrs and "مریم احمدی" in hdrs, "نام کارمندان = ستون‌ها: %s" % hdrs)
wss = wb2["جمع سال"]
found_month_row = any(wss.cell(row=rr, column=1).value == "مرداد" for rr in range(5, 12))
check(found_month_row, "سطر مرداد در شیت جمع سال")
QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: ("", ""))

print("== 5) پشتیبان‌گیری و بازگردانی ==")
bc = win.backup_core
full = bc.snapshot_full()
check(full["version"] == "1.8.1" and full["app"] == "tpp-salary-python", "ساختار بکاپ: version/app")
check(len(full["profiles"]) >= 2 and len(full["records"]) >= 3, "بکاپ شامل کارمندان و رکوردها")
bj = bc.create_json()
bz = bc.create_zip()
check(os.path.exists(bj) and os.path.exists(bz), "فایل JSON و ZIP ساخته شد")
data = bc.load_backup(bz)
check(bc.describe(data)["records"] == len(full["records"]), "خواندن ZIP: تعداد رکورد درست")
check(bc.describe(data)["kind"] == "full", "ZIP شامل full.json")

# تغییر وضعیت → بازگردانی → وضعیت باید به بکاپ برگردد
win.store.save_center(0, "مرکز اضافی بعد از بکاپ")
win.store.upsert_record(e1, cen_id, 1404, 7, {"base_salary": 999, "gross": 999})
recs_now = int(win.db.one("SELECT COUNT(*) AS c FROM records")["c"])
check(recs_now > len(full["records"]), "پس از تغییر، رکورد بیشتری هست (%s)" % recs_now)
summary = bc.restore(data)
check(summary["records"] == len(full["records"]), "بازگردانی: تعداد رکورد = بکاپ (%s)" % summary["records"])
recs_after = int(win.db.one("SELECT COUNT(*) AS c FROM records")["c"])
check(recs_after == len(full["records"]), "رکوردهای اضافی حذف شدند (%s)" % recs_after)
rec1404_7 = win.db.one("SELECT COUNT(*) AS c FROM records WHERE jyear=1404 AND jmonth=7")["c"]
check(int(rec1404_7) == 0, "رکورد ماه ۷ (بعد از بکاپ) پاک شد")
outbox_n = int(win.db.one("SELECT COUNT(*) AS c FROM outbox")["c"])
check(outbox_n == 0, "صف ارسال پس از بازگردانی خالی است")
pend = int(win.db.one("SELECT COUNT(*) AS c FROM employees WHERE pending=1")["c"])
check(pend == 0, "کارمندان پس از بازگردانی pending=0")
emp_back = win.store.employee(e1)
check(emp_back and emp_back["name"] == "علی رضایی", "کارمند بازگردانی‌شده سالم")
# تطبیق با کد ملی: کارمند جدیدی با همان کد ملی بساز → بازگردانی باید همان را تطبیق دهد
e_dup = win.store.save_employee({"id": 0, "name": "نام متفاوت", "national": "1111111111"})
summary2 = bc.restore(bc.load_backup(bj))
check(summary2["users_matched"] >= 2 and summary2["users_created"] == 0,
      "کارمندان با کد ملی تطبیق یافتند و هیچ کارمند جدیدی ساخته نشد (%s)" % summary2["users_matched"])
# هم‌سان افزونه: بازگردانی upsert است — کارمند اضافیِ محلی حذف نمی‌شود
n_emps = int(win.db.one("SELECT COUNT(*) AS c FROM employees WHERE local_deleted=0")["c"])
check(n_emps == len(full["profiles"]) + 1,
      "پروفایل‌های بکاپ اعمال و کارمند اضافی محلی حفظ شد (%s)" % n_emps)

print("== 5b) بازگردانی بکاپ افزونه (سازگاری قالب) ==")
plugin_backup = {
    "version": "1.7.9", "stamp": "2025-01-01 10:00:00",
    "site": {"url": "https://tpptc.ir", "stamp": "2025-01-01 10:00:00"},
    "settings": {"company_name": "شرکت افزونه", "currency": "تومان", "formulas": {}, "defaults": {}},
    "centers": [{"id": 200, "name": "مرکز افزونه", "created_at": "2024-01-01 00:00:00"}],
    "banks": [{"id": 5, "name": "بانک ملت", "sort_order": 0}],
    "fields": [{"field_key": "base_salary", "label": "حقوق پایه", "field_type": "number",
                "default_value": "0", "formula": "", "calculated": 0, "allow_manual": 1, "sort": 1}],
    "records": [{"id": 900, "user_id": 77, "center_id": 200, "jyear": 1404, "jmonth": 8,
                 "payload": "{\"base_salary\": 123000000}", "gross": 123000000, "insurable": 123000000,
                 "insurance_deduct": 9319000, "other_deductions": 0, "net": 113681000,
                 "updated_at": "2025-01-01 10:00:00", "created_at": "2025-01-01 10:00:00"}],
    "profiles": [{"user_id": 77, "user_login": "emp77", "display_name": "کارمند افزونه",
                  "user_email": "emp77@x.ir", "national_id": "999888777", "mobile": "09120000000",
                  "profile": {"full_name": "کارمند افزونه", "centers": [200]}}],
}
summary3 = bc.restore(plugin_backup)
check(summary3["users_created"] == 1 and summary3["records"] == 1, "بکاپ افزونه: کارمند و رکورد جدید")
check(summary3["centers"] == 1 and summary3["banks"] == 1 and summary3["fields"] >= 1, "بکاپ افزونه: بخش‌ها")
emp77 = win.db.one("SELECT * FROM employees WHERE national = '999888777'")
check(emp77 is not None and emp77["name"] == "کارمند افزونه", "کارمند افزونه با کد ملی درج شد")
rec77 = win.db.one("SELECT * FROM records WHERE user_id = ?", (emp77["id"],)) if emp77 else None
check(rec77 is not None, "رکورد افزونه به کارمند نگاشت شد")
if rec77:
    pl = load_json(rec77["payload"], {})
    check(pl.get("base_salary") == 123000000, "پیلود رشته افزونه پارس شد")
    check(rec77["gross"] == 123000000, "مبالغ ستونی رکورد افزونه")
kv77 = win.db.kv_get("company_name", "")
check(kv77 == "شرکت افزونه", "تنظیمات از بکاپ افزونه (company_name)")
# بازگردانی حالت خودی دوباره (وضعیت پاک نشود)
summary4 = bc.restore(bc.load_backup(bj))

print("== 6) صفحات و پنجره ==")
check(len(win.pages) == 12, "۱۲ صفحه در ناوبری: %s" % sorted(win.pages.keys()))
for key in ("payslips", "bankfiche", "annual", "backup"):
    check(key in win.pages, "صفحه جدید «%s» در پنجره" % key)
rep = win.pages["reports"]
REP_YEARS = __import__("app.ui.pages.reports", fromlist=["YEARS"]).YEARS
rep.cmb_year.setCurrentIndex(REP_YEARS.index(1404))
rep.cmb_month.setCurrentIndex(4)
rep.cmb_center.setCurrentIndex(rep.cmb_center.findData(cen_id))
rep.build()
check(len(rep._rows) >= 2, "گزارش لیست حقوق ساخته شد (%s سطر)" % len(rep._rows))
rep_pdf = os.path.join(tmp, "rep_page.pdf")
QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (rep_pdf, ""))
rep.export_pdf()
check(os.path.exists(rep_pdf) and os.path.getsize(rep_pdf) > 3000, "خروجی PDF صفحه گزارش (موتور جدید)")
QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: ("", ""))

win.close()
print()
if fail:
    print("PYAPP-180: FAIL")
    sys.exit(1)
print("PYAPP-180: ALL PASS")
