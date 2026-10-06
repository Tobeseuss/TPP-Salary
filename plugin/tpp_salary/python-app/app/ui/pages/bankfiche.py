# -*- coding: utf-8 -*-
"""فیش بانکی — لیست واریز بانک یک دوره (نسخه 1.8.0).

هم‌سان با صفحه «فیش بانکی» افزونه: کارکنانی که در دوره انتخابی برایشان
حقوق ثبت شده و در بانک انتخابی شماره حساب دارند، به همراه شبا و خالص
پرداختی — با خروجی اکسل و PDF (قالب آینه build_bank_pdf افزونه).
حساب‌ها از پروفایل کارمند (bank_accounts) خوانده می‌شود که با همگام‌سازی
از سایت می‌آید.
"""

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QComboBox, QFileDialog, QFrame, QHBoxLayout, QHeaderView, QLabel,
    QMessageBox, QPushButton, QTableWidget, QTableWidgetItem, QVBoxLayout,
)

from ... import jalali as J
from ... import pdf_engine as PE


class BankFichePage(QFrame):

    def __init__(self, win):
        super().__init__()
        self.win = win
        self._rows = []  # (name, account, sheba, net)

        v = QVBoxLayout(self)
        v.setContentsMargins(22, 20, 22, 20)
        v.setSpacing(12)

        title = QLabel("فیش بانکی")
        title.setObjectName("PageTitle")
        v.addWidget(title)

        hint = QLabel("فهرست واریز حقوق کارکنان دارای حساب در بانک انتخابی — با خروجی اکسل و PDF. شماره حساب‌ها از پروفایل کارمندان (همگام‌شده از سایت) خوانده می‌شود.")
        hint.setObjectName("PageHint")
        hint.setWordWrap(True)
        v.addWidget(hint)

        bar = QHBoxLayout()
        jy, jm, _ = J.today_jalali()
        self.cmb_year = QComboBox()
        for y in range(1399, 1407):
            self.cmb_year.addItem(J.fa_digits(y), y)
        self.cmb_year.setCurrentIndex(max(0, list(range(1399, 1407)).index(jy) if jy in range(1399, 1407) else 6))
        self.cmb_month = QComboBox()
        for m in range(1, 13):
            self.cmb_month.addItem(J.month_name(m), m)
        self.cmb_month.setCurrentIndex(jm - 1)
        self.cmb_center = QComboBox()
        self.cmb_bank = QComboBox()
        btn_show = QPushButton("نمایش فیش بانکی")
        btn_show.setObjectName("Primary")
        btn_show.setCursor(Qt.PointingHandCursor)
        btn_show.clicked.connect(self.load_rows)
        bar.addWidget(self.cmb_year)
        bar.addWidget(self.cmb_month)
        bar.addWidget(self.cmb_center, 1)
        bar.addWidget(self.cmb_bank, 1)
        bar.addWidget(btn_show)
        v.addLayout(bar)

        self.meta_label = QLabel("")
        self.meta_label.setObjectName("PageHint")
        self.meta_label.setWordWrap(True)
        v.addWidget(self.meta_label)

        self.table = QTableWidget(0, 5)
        self.table.setHorizontalHeaderLabels(["#", "نام کارمند", "شماره حساب", "شماره شبا", "حقوق خالص دریافتی"])
        self.table.verticalHeader().hide()
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QTableWidget.SelectRows)
        h = self.table.horizontalHeader()
        h.setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        h.setSectionResizeMode(0, QHeaderView.ResizeMode.Fixed)
        self.table.setColumnWidth(0, 40)
        v.addWidget(self.table, 1)

        foot = QHBoxLayout()
        btn_excel = QPushButton("خروجی اکسل")
        btn_excel.setCursor(Qt.PointingHandCursor)
        btn_excel.clicked.connect(self.export_excel)
        btn_pdf = QPushButton("خروجی PDF")
        btn_pdf.setObjectName("Primary")
        btn_pdf.setCursor(Qt.PointingHandCursor)
        btn_pdf.clicked.connect(self.export_pdf)
        foot.addStretch(1)
        foot.addWidget(btn_excel)
        foot.addWidget(btn_pdf)
        v.addLayout(foot)

    def refresh(self):
        cur_c = self.cmb_center.currentData() or 0
        cur_b = self.cmb_bank.currentData() or 0
        self.cmb_center.blockSignals(True)
        self.cmb_bank.blockSignals(True)
        self.cmb_center.clear()
        self.cmb_center.addItem("همه مراکز", 0)
        for c in self.win.store.centers():
            self.cmb_center.addItem(c["name"], c["id"])
        idx = self.cmb_center.findData(cur_c)
        self.cmb_center.setCurrentIndex(max(0, idx))
        self.cmb_bank.clear()
        self.cmb_bank.addItem("— نام بانک —", 0)
        for b in self.win.store.banks():
            self.cmb_bank.addItem(b["name"], b["id"])
        idx = self.cmb_bank.findData(cur_b)
        self.cmb_bank.setCurrentIndex(max(0, idx))
        self.cmb_center.blockSignals(False)
        self.cmb_bank.blockSignals(False)

    # ---------------- داده ----------------

    def load_rows(self):
        bank_id = self.cmb_bank.currentData() or 0
        if not bank_id:
            QMessageBox.information(self, "فیش بانکی", "ابتدا نام بانک را انتخاب کنید.")
            return
        jyear = self.cmb_year.currentData() or 0
        jmonth = self.cmb_month.currentIndex() + 1
        center_id = self.cmb_center.currentData() or 0
        records = self.win.store.period_records(jyear, jmonth, center_id)
        rows = []
        for rec in records:
            emp = self.win.store.employee(rec["user_id"])
            if not emp:
                continue
            accs = (emp.get("profile") or {}).get("bank_accounts", {}) or {}
            acc = accs.get(str(bank_id)) or accs.get(int(bank_id)) or accs.get(bank_id)
            if not acc or not str(acc.get("account") or "").strip():
                continue
            rows.append((
                str(emp.get("name") or "?"),
                str(acc.get("account") or ""),
                str(acc.get("sheba") or ""),
                float(rec.get("net") or 0),
            ))
        self._rows = rows
        total = sum(r[3] for r in rows)
        bank_name = self.cmb_bank.currentText()
        self.meta_label.setText(
            "کارکنانی که در %s برایشان حقوق ثبت شده و در بانک «%s» شماره حساب دارند: %s نفر — جمع خالص: %s" % (
                J.period_label(jyear, jmonth), bank_name, J.fa_digits(len(rows)), J.format_money(total)))

        self.table.setRowCount(len(rows))
        for r, (name, account, sheba, net) in enumerate(rows):
            it = QTableWidgetItem(J.fa_digits(r + 1))
            it.setTextAlignment(Qt.AlignCenter)
            self.table.setItem(r, 0, it)
            self.table.setItem(r, 1, QTableWidgetItem(name))
            for c, v in ((2, account), (3, sheba)):
                it2 = QTableWidgetItem(v)
                it2.setTextAlignment(Qt.AlignCenter)
                self.table.setItem(r, c, it2)
            it4 = QTableWidgetItem(J.format_money(net))
            it4.setTextAlignment(Qt.AlignCenter)
            self.table.setItem(r, 4, it4)

    # ---------------- خروجی‌ها ----------------

    def _meta(self):
        jyear = self.cmb_year.currentData() or 0
        jmonth = self.cmb_month.currentIndex() + 1
        center_label = self.cmb_center.currentText() if (self.cmb_center.currentData() or 0) else "همه مراکز"
        return jyear, jmonth, center_label

    def export_excel(self):
        if not self._rows:
            QMessageBox.information(self, "فیش بانکی", "ابتدا فیش بانکی را نمایش دهید.")
            return
        try:
            from openpyxl import Workbook
        except ImportError:
            QMessageBox.warning(self, "خروجی", "کتابخانه openpyxl نصب نیست؛ برنامه را دوباره باز کنید تا خودکار نصب شود.")
            return
        jyear, jmonth, center_label = self._meta()
        default = "bank-slip-%s-%s-%02d.xlsx" % (
            self.cmb_bank.currentText(), jyear, jmonth)
        path, _ = QFileDialog.getSaveFileName(self, "ذخیره اکسل فیش بانکی", default, "Excel (*.xlsx)")
        if not path:
            return
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        wb = Workbook()
        ws = wb.active
        ws.title = "فیش بانکی"
        ws.sheet_view.rightToLeft = True
        thin = Side(style="thin", color="B8BFCB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal="center", vertical="center")
        title_font = Font(bold=True, size=12)
        head_font = Font(bold=True, size=10)
        head_fill = PatternFill("solid", fgColor="E8EDF5")
        total_fill = PatternFill("solid", fgColor="F2F2F2")
        numfmt = "#,##0;[Red]-#,##0"

        company = self.win.store.kv_company() or ""
        ws.cell(row=1, column=1, value=("%s — " % company if company else "") + "فیش بانکی %s — %s — %s" % (
            self.cmb_bank.currentText(), J.period_label(jyear, jmonth), center_label))
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
        ws.cell(row=1, column=1).font = title_font
        headers = ["#", "نام کارمند", "شماره حساب", "شماره شبا", "حقوق خالص دریافتی"]
        for c, htxt in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=c, value=htxt)
            cell.font = head_font
            cell.fill = head_fill
            cell.border = border
            cell.alignment = center
        r = 4
        for i, (name, account, sheba, net) in enumerate(self._rows, start=1):
            vals = [i, name, account, sheba, float(net)]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.border = border
                cell.alignment = center
                if c == 5:
                    cell.number_format = numfmt
            r += 1
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        cell = ws.cell(row=r, column=1, value="جمع کل")
        cell.font = head_font
        cell.fill = total_fill
        cell.alignment = center
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = border
            ws.cell(row=r, column=c).fill = total_fill
        cell = ws.cell(row=r, column=5, value=float(sum(x[3] for x in self._rows)))
        cell.font = head_font
        cell.number_format = numfmt
        cell.alignment = center
        for col, w in zip("ABCDE", (6, 32, 24, 28, 22)):
            ws.column_dimensions[col].width = w
        try:
            wb.save(path)
            QMessageBox.information(self, "خروجی اکسل", "فایل ذخیره شد:\n%s" % path)
        except Exception as exc:  # noqa: BLE001
            QMessageBox.warning(self, "خروجی اکسل", "ذخیره ناموفق: %s" % str(exc)[:120])

    def export_pdf(self):
        if not self._rows:
            QMessageBox.information(self, "فیش بانکی", "ابتدا فیش بانکی را نمایش دهید.")
            return
        jyear, jmonth, center_label = self._meta()
        default = "bank-slip-%s-%s-%02d.pdf" % (self.cmb_bank.currentText(), jyear, jmonth)
        path, _ = QFileDialog.getSaveFileName(self, "ذخیره PDF فیش بانکی", default, "PDF (*.pdf)")
        if not path:
            return
        company = self.win.store.kv_company() or ""
        title = ("%s — " % company if company else "") + "فیش بانکی %s — %s — مرکز %s" % (
            self.cmb_bank.currentText(), J.period_label(jyear, jmonth), center_label)
        total = sum(r[3] for r in self._rows)
        meta = ["کارکنان دارای حساب: %s نفر — جمع خالص: %s %s" % (
            J.fa_digits(len(self._rows)), J.format_money(total), self.win.store.kv_currency() or "ریال")]
        try:
            PE.write_bank_pdf(path, self._rows, title, meta, currency=self.win.store.kv_currency() or "ریال")
            QMessageBox.information(self, "خروجی PDF", "فایل ذخیره شد:\n%s" % path)
        except Exception as exc:  # noqa: BLE001
            QMessageBox.warning(self, "خروجی PDF", "ساخت PDF ناموفق: %s" % str(exc)[:160])
