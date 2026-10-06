# -*- coding: utf-8 -*-
"""گزارش سالانه مراکز — نسخه 1.8.0.

گزارش لیست حقوق «یک مرکز» در ماه‌های مختلف «یک سال»؛ خروجی اکسل چندشیتی:
شیت «جمع سال» + به‌ازای هر ماه دارای رکورد یک شیت جداگانه با همان قالب
ستونی 1.6.1 (ستون = نام کارمند، سطر = عناوین حقوق) — کاملاً هم‌سان با
گزارش سالانه افزونه.
"""

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QComboBox, QFileDialog, QFrame, QHBoxLayout, QHeaderView, QLabel,
    QMessageBox, QPushButton, QTableWidget, QTableWidgetItem, QVBoxLayout,
)

from ... import jalali as J
from ...api_client import load_json

YEARS = list(range(1399, 1407))
CALC_ONLY_KEYS = ("overtime_hours", "holiday_days", "absence_days")
SUM_COLS = ["gross", "insurable", "insurance_deduct", "other_deductions", "net"]


def _to_float(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


class AnnualPage(QFrame):

    def __init__(self, win):
        super().__init__()
        self.win = win
        self._stats = []
        self._center = None

        v = QVBoxLayout(self)
        v.setContentsMargins(22, 20, 22, 20)
        v.setSpacing(12)

        title = QLabel("گزارش سالانه مراکز")
        title.setObjectName("PageTitle")
        v.addWidget(title)

        hint = QLabel("گزارش لیست حقوق «یک مرکز» در ماه‌های مختلف «یک سال» — خروجی اکسل چندشیتی: به‌ازای هر ماه دارای رکورد یک شیت جداگانه در کنار شیت «جمع سال» ساخته می‌شود (هم‌سان با افزونه).")
        hint.setObjectName("PageHint")
        hint.setWordWrap(True)
        v.addWidget(hint)

        bar = QHBoxLayout()
        jy, _jm, _jd = J.today_jalali()
        self.cmb_year = QComboBox()
        self.cmb_year.addItems([J.fa_digits(y) for y in YEARS])
        self.cmb_year.setCurrentIndex(max(0, YEARS.index(jy) if jy in YEARS else len(YEARS) - 1))
        self.cmb_center = QComboBox()
        btn_show = QPushButton("نمایش گزارش سالانه")
        btn_show.setObjectName("Primary")
        btn_show.setCursor(Qt.PointingHandCursor)
        btn_show.clicked.connect(self.build)
        bar.addWidget(self.cmb_year)
        bar.addWidget(self.cmb_center, 1)
        bar.addWidget(btn_show)
        v.addLayout(bar)

        self.meta_label = QLabel("")
        self.meta_label.setObjectName("PageHint")
        self.meta_label.setWordWrap(True)
        v.addWidget(self.meta_label)

        self.table = QTableWidget(0, 7)
        self.table.setHorizontalHeaderLabels(
            ["ماه", "تعداد فیش", "جمع ناخالص", "جمع مشمول بیمه", "بیمه سهم کارمند", "سایر کسورات", "جمع خالص پرداختی"])
        self.table.verticalHeader().hide()
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        h = self.table.horizontalHeader()
        h.setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        v.addWidget(self.table, 1)

        foot = QHBoxLayout()
        btn_excel = QPushButton("خروجی اکسل سالانه (چندشیتی)")
        btn_excel.setObjectName("Primary")
        btn_excel.setCursor(Qt.PointingHandCursor)
        btn_excel.clicked.connect(self.export_excel)
        foot.addStretch(1)
        foot.addWidget(btn_excel)
        v.addLayout(foot)

        self._stats = []

    def refresh(self):
        cur = self.cmb_center.currentData() or 0
        self.cmb_center.blockSignals(True)
        self.cmb_center.clear()
        self.cmb_center.addItem("— مرکز —", 0)
        for c in self.win.store.centers():
            self.cmb_center.addItem(c["name"], c["id"])
        idx = self.cmb_center.findData(cur)
        self.cmb_center.setCurrentIndex(max(0, idx))
        self.cmb_center.blockSignals(False)

    # ---------------- ساخت گزارش ----------------

    def build(self):
        center_id = self.cmb_center.currentData() or 0
        if not center_id:
            QMessageBox.information(self, "گزارش سالانه", "ابتدا مرکز را انتخاب کنید.")
            return
        jyear = YEARS[self.cmb_year.currentIndex()]
        center = self.win.store.center(center_id)
        self._center = dict(center) if center else None

        stats = []
        for m in range(1, 13):
            recs = self.win.store.period_records(jyear, m, center_id)
            if not recs:
                continue
            row = {"jmonth": m, "count": len(recs)}
            for k in SUM_COLS:
                row[k] = sum(_to_float(r.get(k)) for r in recs)
            stats.append(row)
        self._stats = stats

        self.table.setRowCount(len(stats) + (1 if stats else 0))
        for r, s in enumerate(stats):
            vals = [J.month_name(s["jmonth"]), J.fa_digits(s["count"])] + \
                   [J.format_money(s[k]) for k in SUM_COLS]
            for c, v in enumerate(vals):
                it = QTableWidgetItem(str(v))
                if c >= 1:
                    it.setTextAlignment(Qt.AlignCenter)
                if c == 6:
                    f = it.font()
                    f.setBold(True)
                    it.setFont(f)
                self.table.setItem(r, c, it)
        if stats:
            total = {k: sum(s[k] for s in stats) for k in SUM_COLS}
            tvals = ["جمع سال", J.fa_digits(sum(s["count"] for s in stats))] + \
                    [J.format_money(total[k]) for k in SUM_COLS]
            for c, v in enumerate(tvals):
                it = QTableWidgetItem(str(v))
                it.setTextAlignment(Qt.AlignCenter if c else Qt.AlignRight)
                f = it.font()
                f.setBold(True)
                it.setFont(f)
                it.setBackground(Qt.lightGray)
                self.table.setItem(len(stats), c, it)
        self.meta_label.setText(
            "سال: %s — مرکز: %s — %s ماه دارای رکورد" % (
                J.fa_digits(jyear), (self._center or {}).get("name", "—"), J.fa_digits(len(stats))))

    # ---------------- خروجی اکسل چندشیتی ----------------

    def export_excel(self):
        if not self._stats:
            QMessageBox.information(self, "گزارش سالانه", "ابتدا گزارش را بسازید.")
            return
        try:
            from openpyxl import Workbook
        except ImportError:
            QMessageBox.warning(self, "خروجی", "کتابخانه openpyxl نصب نیست؛ برنامه را دوباره باز کنید تا خودکار نصب شود.")
            return
        jyear = YEARS[self.cmb_year.currentIndex()]
        center_id = self.cmb_center.currentData() or 0
        center_name = (self._center or {}).get("name", "مرکز")
        company = self.win.store.kv_company() or ""
        currency = self.win.store.kv_currency() or "ریال"

        default = "annual-%s-%s.xlsx" % (jyear, str(center_name).strip()[:20])
        path, _ = QFileDialog.getSaveFileName(self, "ذخیره اکسل سالانه", default, "Excel (*.xlsx)")
        if not path:
            return

        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        title_font = Font(bold=True, size=12)
        head_font = Font(bold=True, size=10)
        head_fill = PatternFill("solid", fgColor="E8EDF5")
        label_fill = PatternFill("solid", fgColor="F7F7F7")
        total_fill = PatternFill("solid", fgColor="F2F2F2")
        thin = Side(style="thin", color="B8BFCB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center_a = Alignment(horizontal="center", vertical="center")
        right_a = Alignment(horizontal="right", vertical="center")
        numfmt = "#,##0;[Red]-#,##0"

        # --- شیت جمع سال ---
        ws = wb.active
        ws.title = "جمع سال"
        ws.sheet_view.rightToLeft = True
        ws.cell(row=1, column=1, value=("%s — " % company if company else "") +
                "گزارش سالانه %s — مرکز %s" % (J.fa_digits(jyear), center_name)).font = title_font
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
        ws.cell(row=2, column=1, value="واحد: %s" % currency).font = head_font
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=7)
        heads = ["ماه", "تعداد فیش", "جمع ناخالص", "جمع مشمول بیمه", "بیمه سهم کارمند", "سایر کسورات", "جمع خالص پرداختی"]
        for c, htxt in enumerate(heads, start=1):
            cell = ws.cell(row=4, column=c, value=htxt)
            cell.font = head_font
            cell.fill = head_fill
            cell.border = border
            cell.alignment = center_a
        r = 5
        for s in self._stats:
            vals = [J.month_name(s["jmonth"]), s["count"]] + [float(s[k]) for k in SUM_COLS]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.border = border
                cell.alignment = center_a
                if c >= 3:
                    cell.number_format = numfmt
                if c == 2:
                    cell.number_format = "0"
            r += 1
        if self._stats:
            total = {k: sum(s[k] for s in self._stats) for k in SUM_COLS}
            vals = ["جمع سال", sum(s["count"] for s in self._stats)] + [float(total[k]) for k in SUM_COLS]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.font = head_font
                cell.fill = total_fill
                cell.border = border
                cell.alignment = center_a
                if c >= 3:
                    cell.number_format = numfmt
            r += 1
        for col, w in zip("ABCDEFG", (12, 12, 20, 20, 18, 18, 22)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A5"

        # --- شیت هر ماه (قالب ستونی 1.6.1) ---
        emp_name = {}
        fields = self.win.store.fields()
        for s in self._stats:
            recs = self.win.store.period_records(jyear, s["jmonth"], center_id)
            if not recs:
                continue
            payloads = [load_json(rec.get("payload"), {}) or {} for rec in recs]
            # فیلدهای چاپی — آینه printable_fields افزونه
            printable = []
            for f in fields:
                key = str(f.get("key", ""))
                if key in CALC_ONLY_KEYS:
                    continue
                if str(f.get("type", "")) == "number":
                    if not any(abs(_to_float(p.get(key))) > 1e-4 for p in payloads):
                        continue
                printable.append((key, str(f.get("label") or key), str(f.get("type", ""))))

            used_names = set()
            for rec in recs:
                uid = int(rec["user_id"])
                if uid not in emp_name:
                    emp = self.win.store.employee(uid)
                    emp_name[uid] = emp["name"] if emp else "کارمند #%s" % uid
                used_names.add(uid)

            sheet_name = J.month_name(s["jmonth"])[:31]
            ws2 = wb.create_sheet(sheet_name)
            ws2.sheet_view.rightToLeft = True
            ncols = 1 + len(recs)
            ws2.cell(row=1, column=1, value=("%s — " % company if company else "") +
                     "لیست حقوق %s %s — %s" % (J.month_name(s["jmonth"]), J.fa_digits(jyear), center_name)).font = title_font
            ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
            ws2.cell(row=2, column=1, value="واحد: %s" % currency).font = head_font
            ws2.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
            for c in range(1, ncols + 1):
                cell = ws2.cell(row=4, column=c)
                cell.font = head_font
                cell.fill = head_fill
                cell.border = border
                cell.alignment = center_a
            ws2.cell(row=4, column=1, value="عناوین")
            for k, rec in enumerate(recs):
                ws2.cell(row=4, column=2 + k, value=emp_name[int(rec["user_id"])])
            r2 = 5
            for key, label, ftype in printable:
                cell = ws2.cell(row=r2, column=1, value=label)
                cell.font = head_font
                cell.fill = label_fill
                cell.border = border
                cell.alignment = right_a
                for k, rec in enumerate(recs):
                    payload = load_json(rec.get("payload"), {}) or {}
                    cell = ws2.cell(row=r2, column=2 + k)
                    cell.border = border
                    cell.alignment = center_a
                    if ftype == "number":
                        cell.value = _to_float(payload.get(key))
                        cell.number_format = numfmt
                    else:
                        cell.value = str(payload.get(key, "") or "")
                r2 += 1
            ws2.column_dimensions["A"].width = 26
            for c in range(2, ncols + 1):
                ws2.column_dimensions[get_column_letter(c)].width = 22
            ws2.freeze_panes = "A5"

        try:
            wb.save(path)
            QMessageBox.information(self, "خروجی اکسل سالانه", "فایل ذخیره شد:\n%s" % path)
        except Exception as exc:  # noqa: BLE001
            QMessageBox.warning(self, "خروجی اکسل سالانه", "ذخیره ناموفق: %s" % str(exc)[:120])
